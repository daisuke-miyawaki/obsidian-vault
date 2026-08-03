#!/usr/bin/env python3
"""マーケティング・先読みサイクル年表 Excelテンプレート生成"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

OUTPUT = Path(__file__).parent / "マーケティング・先読みサイクル年表.xlsx"


def create_data_sheet(wb: Workbook):
    ws = wb.active
    ws.title = "データ"

    headers = [
        "項目",
        "US_導入開始",
        "US_成長開始",
        "US_成熟開始",
        "US_衰退開始",
        "US_終了",
        "JP_導入開始",
        "JP_成長開始",
        "JP_成熟開始",
        "JP_衰退開始",
        "JP_終了",
        "us_hint",
        "jp_hint",
    ]
    ws.append(headers)

    ws.append(
        [
            "LP/ファネル",
            "2008-01",
            "2012-01",
            "2016-01",
            "2022-01",
            "2030-12",
            "2012-01",
            "2016-01",
            "2020-01",
            "2027-01",
            "2032-12",
            "",
            "",
        ]
    )
    ws.append(
        [
            "LP/ファネル",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "米国のLP/ファネルはダイレクトレスポンス文化と短い導線に支えられ普及したが、現在は成熟期から衰退期へ移行中。",
            "日本は検討期間が長く、機械的な量産コンテンツへの警戒感から米国より展開が遅れる。",
        ]
    )

    for col in range(1, len(headers) + 1):
        ws.cell(1, col).font = Font(bold=True)
    ws.column_dimensions["A"].width = 14
    for col in "BCDEFGHIJK":
        ws.column_dimensions[col].width = 12
    ws.column_dimensions["L"].width = 40
    ws.column_dimensions["M"].width = 40

    note = ws.cell(5, 1)
    note.value = (
        "※ GemのCSV（2行/項目）を2行目から貼り付け。"
        "貼ったらターミナルで python3 refresh_bands.py を実行し、「年表」シートを開き直す。"
    )
    note.font = Font(italic=True, color="666666")


def main():
    wb = Workbook()
    create_data_sheet(wb)
    wb.create_sheet("年表")
    wb.save(OUTPUT)
    print(f"Created: {OUTPUT}")

    from refresh_bands import refresh

    refresh()


if __name__ == "__main__":
    main()
