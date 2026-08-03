#!/usr/bin/env python3
"""データシートの内容を読み、年表シートに太い横棒（4期）を描画する。"""

from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

XLSX = Path(__file__).parent / "マーケティング・先読みサイクル年表.xlsx"

PHASE_FILLS = {
    1: PatternFill("solid", fgColor="DCE6F1"),  # 導入期
    2: PatternFill("solid", fgColor="A9D18E"),  # 成長期
    3: PatternFill("solid", fgColor="FFE699"),  # 成熟期
    4: PatternFill("solid", fgColor="BFBFBF"),  # 衰退期
}
PHASE_LABELS = {1: "導入", 2: "成長", 3: "成熟", 4: "衰退"}
BLOCK_FILL = PatternFill("solid", fgColor="404040")
BLOCK_FONT = Font(bold=True, color="FFFFFF", size=11)
HINT_FILL = PatternFill("solid", fgColor="F2F2F2")
EMPTY_FILL = PatternFill(fill_type=None)
THIN = Side(style="thin", color="CCCCCC")

YEAR_PAD = 1  # データの前後に余白年数
BAND_HEIGHT = 28
HINT_HEIGHT = 44
HEADER_ROW = 4
FIRST_BLOCK_ROW = 5


def parse_ym(s) -> date | None:
    if not s or not isinstance(s, str) or "-" not in str(s):
        return None
    y, m = str(s).split("-", 1)
    return date(int(y), int(m), 1)


def phase_at(dt: date, bounds) -> int:
    intro, growth, mature, decline, end = [parse_ym(b) for b in bounds]
    if not all([intro, growth, mature, decline, end]):
        return 0
    if dt < intro or dt > end:
        return 0
    if dt < growth:
        return 1
    if dt < mature:
        return 2
    if dt < decline:
        return 3
    return 4


def phase_segments(bounds, year_start: int, year_end: int) -> list[tuple[int, int, int]]:
    """連続する同じ期の (phase, start_year, end_year) のリスト。"""
    segments: list[tuple[int, int, int]] = []
    current_phase = None
    seg_start = None

    for year in range(year_start, year_end + 1):
        p = phase_at(date(year, 7, 1), bounds)
        if p == 0:
            if current_phase is not None:
                segments.append((current_phase, seg_start, year - 1))
                current_phase = None
                seg_start = None
            continue
        if p != current_phase:
            if current_phase is not None:
                segments.append((current_phase, seg_start, year - 1))
            current_phase = p
            seg_start = year

    if current_phase is not None:
        segments.append((current_phase, seg_start, year_end))

    return segments


def read_items(data_ws):
    items = []
    row = 2
    while row <= data_ws.max_row:
        name = data_ws.cell(row, 1).value
        if not name or str(name).startswith("※"):
            break
        us = [data_ws.cell(row, c).value for c in range(2, 7)]
        jp = [data_ws.cell(row, c).value for c in range(7, 12)]
        hint_row = row + 1
        us_hint = data_ws.cell(hint_row, 12).value or ""
        jp_hint = data_ws.cell(hint_row, 13).value or ""
        if any(us) or any(jp):
            items.append(
                {
                    "name": str(name),
                    "us": us,
                    "jp": jp,
                    "us_hint": str(us_hint),
                    "jp_hint": str(jp_hint),
                }
            )
        row += 2
    return items


def year_range(items) -> tuple[int, int]:
    years = []
    for item in items:
        for bounds in (item["us"], item["jp"]):
            intro = parse_ym(bounds[0])
            end = parse_ym(bounds[4])
            if intro:
                years.append(intro.year)
            if end:
                years.append(end.year)
    if not years:
        return 2010, 2035
    return min(years) - YEAR_PAD, max(years) + YEAR_PAD


def year_to_col(year: int, year_start: int) -> int:
    return year - year_start + 2


def clear_chart_sheet(chart_ws, from_row: int, num_year_cols: int):
    for merged in list(chart_ws.merged_cells.ranges):
        if merged.min_row >= from_row:
            chart_ws.unmerge_cells(str(merged))

    for r in range(from_row, chart_ws.max_row + 20):
        for c in range(1, num_year_cols + 2):
            cell = chart_ws.cell(r, c)
            try:
                cell.value = None
            except AttributeError:
                pass
            cell.fill = EMPTY_FILL
            cell.border = Border()


def draw_band_row(ws, row: int, label: str, bounds, year_start: int, year_end: int):
    ws.cell(row, 1, label)
    ws.cell(row, 1).font = Font(bold=True)
    ws.cell(row, 1).alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = BAND_HEIGHT

    for phase, y0, y1 in phase_segments(bounds, year_start, year_end):
        c0 = year_to_col(y0, year_start)
        c1 = year_to_col(y1, year_start)
        if c0 > c1:
            continue
        if c0 != c1:
            ws.merge_cells(start_row=row, start_column=c0, end_row=row, end_column=c1)
        cell = ws.cell(row, c0)
        cell.fill = PHASE_FILLS[phase]
        cell.value = PHASE_LABELS[phase]
        cell.font = Font(size=9, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        for c in range(c0, c1 + 1):
            ws.cell(row, c).border = Border(
                top=THIN, bottom=THIN, left=THIN if c == c0 else THIN, right=THIN if c == c1 else THIN
            )


def draw_block_header(ws, row: int, title: str, year_start: int, year_end: int):
    ws.row_dimensions[row].height = 22
    ws.cell(row, 1, title)
    ws.cell(row, 1).fill = BLOCK_FILL
    ws.cell(row, 1).font = BLOCK_FONT
    c0 = 2
    c1 = year_to_col(year_end, year_start)
    if c1 > c0:
        ws.merge_cells(start_row=row, start_column=c0, end_row=row, end_column=c1)
    cell = ws.cell(row, c0, title)
    cell.fill = BLOCK_FILL
    cell.font = BLOCK_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")


def draw_hint_row(ws, row: int, us_hint: str, jp_hint: str, year_start: int, year_end: int):
    ws.cell(row, 1, "　")
    ws.row_dimensions[row].height = HINT_HEIGHT
    c_mid = year_to_col((year_start + year_end) // 2, year_start)
    c_end = year_to_col(year_end, year_start)

    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=c_mid)
    ws.merge_cells(start_row=row, start_column=c_mid + 1, end_row=row, end_column=c_end)

    uh = ws.cell(row, 2, f"米国：{us_hint}" if us_hint else "")
    jh = ws.cell(row, c_mid + 1, f"日本：{jp_hint}" if jp_hint else "")
    for c in (uh, jh):
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.fill = HINT_FILL
        c.font = Font(size=9, color="444444")


def setup_header(chart_ws, year_start: int, year_end: int):
    chart_ws["A1"] = "マーケティング・先読みサイクル年表"
    chart_ws["A1"].font = Font(bold=True, size=14)

    legend_parts = [
        ("導入期", PHASE_FILLS[1]),
        ("成長期", PHASE_FILLS[2]),
        ("成熟期", PHASE_FILLS[3]),
        ("衰退期", PHASE_FILLS[4]),
    ]
    chart_ws["A2"] = "凡例："
    for i, (label, fill) in enumerate(legend_parts):
        c = chart_ws.cell(2, 2 + i * 2, label)
        c.fill = fill
        c.font = Font(bold=True, size=9)

    chart_ws["A3"] = f"横軸＝年（{year_start}〜{year_end}）｜帯の色＝MBA 4期"
    chart_ws["A3"].font = Font(italic=True, color="666666", size=9)

    chart_ws.cell(HEADER_ROW, 1, "項目")
    chart_ws.cell(HEADER_ROW, 1).font = Font(bold=True)
    chart_ws.column_dimensions["A"].width = 18

    for year in range(year_start, year_end + 1):
        col = year_to_col(year, year_start)
        cell = chart_ws.cell(HEADER_ROW, col, year)
        cell.font = Font(bold=True, size=9)
        cell.alignment = Alignment(horizontal="center")
        chart_ws.column_dimensions[get_column_letter(col)].width = 5.5

    chart_ws.freeze_panes = f"B{FIRST_BLOCK_ROW}"


def refresh():
    wb = load_workbook(XLSX)
    data_ws = wb["データ"]
    chart_ws = wb["年表"]
    items = read_items(data_ws)

    if not items:
        print("No items in データ sheet.")
        return

    year_start, year_end = year_range(items)
    num_cols = year_end - year_start + 1

    clear_chart_sheet(chart_ws, HEADER_ROW, num_cols)
    setup_header(chart_ws, year_start, year_end)

    current_row = FIRST_BLOCK_ROW

    draw_block_header(chart_ws, current_row, "アメリカ", year_start, year_end)
    current_row += 1

    for item in items:
        draw_band_row(chart_ws, current_row, item["name"], item["us"], year_start, year_end)
        current_row += 1

    current_row += 1
    draw_block_header(chart_ws, current_row, "日本", year_start, year_end)
    current_row += 1

    for item in items:
        draw_band_row(chart_ws, current_row, item["name"], item["jp"], year_start, year_end)
        current_row += 1
        draw_hint_row(chart_ws, current_row, item["us_hint"], item["jp_hint"], year_start, year_end)
        current_row += 1

    wb.save(XLSX)
    print(f"Refreshed {len(items)} item(s), years {year_start}-{year_end} → {XLSX}")


if __name__ == "__main__":
    refresh()
